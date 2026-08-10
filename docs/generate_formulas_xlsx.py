"""Generate CSV_Reader_Formulas.xlsx — list of all metrics formulas for verification."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent / "CSV_Reader_Formulas.xlsx"

# Columns: ID, Metric, Category, Formula (as implemented), Plain English, Code location, Status / Notes
FORMULAS = [
    (
        "F01",
        "Total PnL",
        "Core",
        "sum(pnl of all trades)",
        "Sum of every trade's PnL.",
        "metrics.py ~85",
        "OK — standard",
    ),
    (
        "F02",
        "Total Trades",
        "Core",
        "count(trades)",
        "Number of trades after valid entry_time.",
        "metrics.py ~84",
        "OK — standard",
    ),
    (
        "F03",
        "Win Trades",
        "Core",
        "count(pnl > 0)",
        "Trades with strictly positive PnL.",
        "metrics.py ~86-88",
        "OK — breakeven NOT counted as win",
    ),
    (
        "F04",
        "Loss Trades",
        "Core",
        "count(pnl <= 0)",
        "Trades with zero or negative PnL.",
        "metrics.py ~87-89",
        "REVIEW — zero PnL treated as loss",
    ),
    (
        "F05",
        "Win Rate (%)",
        "Core",
        "(win_trades / total_trades) * 100",
        "Percentage of winning trades.",
        "metrics.py ~90",
        "OK — rounded to 3 decimals",
    ),
    (
        "F06",
        "Avg Win",
        "Trade stats",
        "mean(pnl where pnl > 0)",
        "Average PnL of winning trades only.",
        "metrics.py ~92",
        "OK — standard",
    ),
    (
        "F07",
        "Avg Loss",
        "Trade stats",
        "mean(pnl where pnl <= 0)",
        "Average PnL of losing/breakeven trades (usually negative).",
        "metrics.py ~93",
        "REVIEW — includes zeros with losses",
    ),
    (
        "F08",
        "Avg PnL per Trade",
        "Trade stats",
        "total_pnl / total_trades",
        "Mean PnL across all trades.",
        "metrics.py ~140",
        "OK — standard",
    ),
    (
        "F09",
        "Risk Reward",
        "Trade stats",
        "abs(avg_win / avg_loss)",
        "How large wins are vs losses (absolute ratio).",
        "metrics.py ~141-145",
        "OK — returns 0 if avg_loss is 0",
    ),
    (
        "F10",
        "Expectancy",
        "Trade stats",
        "(win_frac * avg_win) + (loss_frac * avg_loss)",
        "Expected PnL per trade from win rate and avg win/loss.",
        "metrics.py ~147-149",
        "OK — classic expectancy",
    ),
    (
        "F11",
        "Gross Wins",
        "Trade stats",
        "sum(pnl where pnl > 0)",
        "Total money made on winners.",
        "metrics.py ~151",
        "OK — intermediate for profit factor",
    ),
    (
        "F12",
        "Gross Losses",
        "Trade stats",
        "abs(sum(pnl where pnl <= 0))",
        "Absolute total of losing/breakeven PnL.",
        "metrics.py ~152",
        "OK — intermediate for profit factor",
    ),
    (
        "F13",
        "Profit Factor",
        "Trade stats",
        "gross_wins / gross_losses",
        "Gross profit divided by gross loss. >1 means net profitable on wins vs losses.",
        "metrics.py ~153-155",
        "OK — classic; 0 if no losses",
    ),
    (
        "F14",
        "Daily PnL",
        "Time series",
        "groupby(trade_date).sum(pnl)",
        "All trades on the same calendar day (by entry date) summed.",
        "metrics.py ~82,95",
        "OK — uses entry_time date, not exit",
    ),
    (
        "F15",
        "Total Days",
        "Time series",
        "count(unique trade_dates)",
        "Number of days that had at least one trade.",
        "metrics.py ~96",
        "OK — trading days only, not calendar span",
    ),
    (
        "F16",
        "Profit Days",
        "Time series",
        "count(daily_pnl > 0)",
        "Days where net daily PnL is positive.",
        "metrics.py ~97",
        "OK — standard",
    ),
    (
        "F17",
        "Loss Days",
        "Time series",
        "count(daily_pnl < 0)",
        "Days where net daily PnL is negative (zero days excluded).",
        "metrics.py ~98",
        "OK — flat days not counted as loss days",
    ),
    (
        "F18",
        "Profit Days %",
        "Time series",
        "(profit_days / total_days) * 100",
        "Share of trading days that were profitable.",
        "metrics.py ~99",
        "OK — standard",
    ),
    (
        "F19",
        "Avg Profit on Profit Days",
        "Time series",
        "mean(daily_pnl where daily_pnl > 0)",
        "Average size of green days.",
        "metrics.py ~101-103",
        "OK — standard",
    ),
    (
        "F20",
        "Avg Loss on Loss Days",
        "Time series",
        "mean(daily_pnl where daily_pnl < 0)",
        "Average size of red days.",
        "metrics.py ~104-106",
        "OK — standard",
    ),
    (
        "F21",
        "Equity Curve",
        "Drawdown",
        "cumsum(daily_pnl)",
        "Running total of daily PnL (starts from first trading day).",
        "metrics.py ~108",
        "OK — no starting capital added",
    ),
    (
        "F22",
        "Drawdown (series)",
        "Drawdown",
        "equity - cummax(equity)",
        "How far equity is below its peak so far (≤ 0).",
        "metrics.py ~109",
        "OK — classic peak-to-trough DD in PnL units",
    ),
    (
        "F23",
        "Max Drawdown",
        "Drawdown",
        "min(drawdown series)",
        "Worst (most negative) drawdown value.",
        "metrics.py ~110",
        "OK — in currency/PnL units, not %",
    ),
    (
        "F24",
        "Top-5 DD Avg",
        "Drawdown",
        "mean(5 most negative drawdown point values)",
        "Average of the 5 worst drawdown *points* (sorted ascending).",
        "metrics.py ~112-115",
        "REVIEW — not distinct DD events; just worst point values",
    ),
    (
        "F25",
        "Top-5 Loss Day Avg",
        "Drawdown",
        "mean(5 most negative daily_pnl values)",
        "Average of the 5 worst single-day losses.",
        "metrics.py ~117-119",
        "OK — clear definition",
    ),
    (
        "F26",
        "Calmar",
        "Risk ratios",
        "abs(total_pnl / max_drawdown)",
        "Total PnL over |max DD|.",
        "metrics.py ~121",
        "REVIEW — classic Calmar uses annualized return / max DD %",
    ),
    (
        "F27",
        "Recovery Factor",
        "Risk ratios",
        "abs(total_pnl / max_drawdown)",
        "Same formula as Calmar in this codebase (different rounding).",
        "metrics.py ~122-124",
        "REVIEW — identical to Calmar; only rounds to 3 vs 2 decimals",
    ),
    (
        "F28",
        "Monthly PnL",
        "Monthly",
        "groupby(year-month).sum(pnl)",
        "Sum of trade PnL by calendar month of entry.",
        "metrics.py ~126-127",
        "OK — standard",
    ),
    (
        "F29",
        "Total Months",
        "Monthly",
        "count(months with ≥1 trade)",
        "Months that had activity (not full calendar span).",
        "metrics.py ~128",
        "OK",
    ),
    (
        "F30",
        "Avg Monthly",
        "Monthly",
        "mean(monthly_pnl)",
        "Average PnL across active months.",
        "metrics.py ~129",
        "OK",
    ),
    (
        "F31",
        "Median Monthly",
        "Monthly",
        "median(monthly_pnl)",
        "Median monthly PnL (robust to outliers).",
        "metrics.py ~130",
        "OK",
    ),
    (
        "F32",
        "Sortino",
        "Risk ratios",
        "avg_monthly / std(negative monthly_pnl only, ddof=0)",
        "Avg monthly return divided by std of losing months only.",
        "metrics.py ~132-134",
        "REVIEW — not classic Sortino (usually downside of all returns below target)",
    ),
    (
        "F33",
        "Avg Weekly",
        "Weekly",
        "mean(groupby(ISO week number).sum(pnl))",
        "Average of weekly PnL buckets.",
        "metrics.py ~136-138",
        "BUG RISK — groups by week number only (1–53), NOT year+week; years merge",
    ),
    (
        "F34",
        "Mean Daily",
        "Volatility",
        "mean(daily_pnl)",
        "Average daily PnL on trading days.",
        "metrics.py ~157",
        "OK",
    ),
    (
        "F35",
        "Std Dev (daily)",
        "Volatility",
        "std(daily_pnl, ddof=0)",
        "Population standard deviation of daily PnL.",
        "metrics.py ~158",
        "OK — population (ddof=0), not sample",
    ),
    (
        "F36",
        "Mean / Std",
        "Volatility",
        "mean_daily / std_dev",
        "Raw signal-to-noise of daily PnL (not annualized).",
        "metrics.py ~159",
        "OK",
    ),
    (
        "F37",
        "Sharpe",
        "Risk ratios",
        "(mean_daily / std_dev) * sqrt(252)",
        "Annualized Sharpe-like ratio on daily PnL (rf=0).",
        "metrics.py ~160-162",
        "REVIEW — uses PnL not returns; assumes 252 trading days; rf=0",
    ),
    (
        "F38",
        "SQN (System Quality Number)",
        "Risk ratios",
        "sqrt(N) * (avg_pnl_per_trade / std(trade_pnl, ddof=0))",
        "Van Tharp SQN-style score from trade-level mean/std.",
        "metrics.py ~163-167",
        "OK — classic shape; needs N>1",
    ),
    (
        "F39",
        "Monthly heatmap cell",
        "UI / Monthly",
        "sum(pnl) for trades in that year+month",
        "Each cell = month PnL; year net = sum of 12 months.",
        "metrics.py ~240-264",
        "OK",
    ),
    (
        "F40",
        "Grand Total (monthly)",
        "UI / Monthly",
        "total_pnl (same as F01)",
        "Footer total equals overall Total PnL.",
        "metrics.py ~264",
        "OK",
    ),
]

TEST_STEPS = [
    (
        1,
        "Unit check with tiny CSV",
        "Make a 5–10 trade CSV with known PnLs. Hand-calculate Total PnL, Win Rate, Profit Factor. Upload to http://localhost:5174 or http://<server>:8080 and compare Overview.",
    ),
    (
        2,
        "Run verify_formulas.py",
        "From csv_reader/:  python docs/verify_formulas.py\nUses samples/sample_trades.csv, recomputes formulas independently, compares to compute_analysis().",
    ),
    (
        3,
        "API check",
        "Upload CSV → GET /api/overview (cookie session). Diff JSON fields vs Excel expected formulas.",
    ),
    (
        4,
        "Equity / DD visual check",
        "On Equity page: equity should be cumulative daily PnL; drawdown should never go above 0; max DD should match the deepest trough.",
    ),
    (
        5,
        "Monthly heatmap check",
        "Pick one month: sum that month's trade PnLs in Excel → must match heatmap cell; year Net must match sum of months.",
    ),
    (
        6,
        "Flag REVIEW / BUG RISK rows",
        "Especially F26/F27 (Calmar=Recovery), F32 (Sortino), F33 (weekly year merge), F37 (Sharpe on PnL). Confirm with team before 'critically improve'.",
    ),
]

HAND_EXAMPLE = [
    ("Trade", "Date", "PnL"),
    (1, "2024-01-02", 100),
    (2, "2024-01-02", -40),
    (3, "2024-01-03", 50),
    (4, "2024-01-04", -20),
    (5, "2024-01-05", 0),
]

HAND_EXPECTED = [
    ("Metric", "Expected value", "How"),
    ("Total PnL", 90, "100-40+50-20+0"),
    ("Total Trades", 5, "count"),
    ("Win Trades", 2, "pnl>0 → trades 1,3"),
    ("Loss Trades", 3, "pnl<=0 → trades 2,4,5"),
    ("Win Rate %", 40.0, "2/5*100"),
    ("Avg Win", 75.0, "(100+50)/2"),
    ("Avg Loss", -20.0, "(-40-20+0)/3"),
    ("Daily PnL days", "60, 50, -20, 0", "Jan2:60, Jan3:50, Jan4:-20, Jan5:0"),
    ("Profit Days", 2, "days with daily>0"),
    ("Loss Days", 1, "days with daily<0 (flat day excluded)"),
    ("Equity end", 90, "cumsum"),
    ("Max DD", "compute from peak", "equity peaks then dips"),
    ("Profit Factor", "150 / 60 = 2.5", "gross wins 150 / abs gross losses 60"),
    ("Expectancy", "0.4*75 + 0.6*(-20) = 18", "win_frac*avg_win + loss_frac*avg_loss"),
]


def _style_header(ws, row=1):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin


def _autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    wb = Workbook()

    # Sheet 1 — All formulas
    ws = wb.active
    ws.title = "All Formulas"
    headers = [
        "ID",
        "Metric (UI / API)",
        "Category",
        "Formula (as coded)",
        "Plain English",
        "Code location",
        "Verification status / notes",
    ]
    ws.append(headers)
    _style_header(ws)

    review_fill = PatternFill("solid", fgColor="FFF2CC")
    bug_fill = PatternFill("solid", fgColor="FCE4D6")
    ok_fill = PatternFill("solid", fgColor="E2EFDA")

    for row in FORMULAS:
        ws.append(list(row))
        r = ws.max_row
        note = row[6]
        fill = ok_fill
        if note.startswith("REVIEW"):
            fill = review_fill
        elif note.startswith("BUG"):
            fill = bug_fill
        for c in range(1, 8):
            cell = ws.cell(r, c)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{ws.max_row}"
    _autosize(ws, [6, 28, 14, 55, 45, 16, 55])
    ws.row_dimensions[1].height = 30

    # Sheet 2 — How to test
    ws2 = wb.create_sheet("How to Test")
    ws2.append(["Step", "Test", "Instructions"])
    _style_header(ws2)
    for row in TEST_STEPS:
        ws2.append(list(row))
        for c in range(1, 4):
            ws2.cell(ws2.max_row, c).alignment = Alignment(wrap_text=True, vertical="top")
    for r in range(2, ws2.max_row + 1):
        ws2.row_dimensions[r].height = 55
    _autosize(ws2, [8, 28, 90])

    # Sheet 3 — Hand example
    ws3 = wb.create_sheet("Hand Example (5 trades)")
    for row in HAND_EXAMPLE:
        ws3.append(list(row))
    _style_header(ws3)
    ws3.append([])
    start = ws3.max_row + 1
    for i, row in enumerate(HAND_EXPECTED):
        ws3.append(list(row))
        if i == 0:
            _style_header(ws3, ws3.max_row)
    _autosize(ws3, [22, 28, 50])

    # Sheet 4 — Source map
    ws4 = wb.create_sheet("Source Map")
    ws4.append(["File", "Role"])
    _style_header(ws4)
    for row in [
        ("csv_reader/backend/app/core/metrics.py", "All formula implementations (single source of truth)"),
        ("csv_reader/backend/app/schemas/analysis.py", "API field names for Overview / Equity / Monthly"),
        ("csv_reader/frontend (Overview page)", "Displays overview metrics from /api/overview"),
        ("csv_reader/docs/verify_formulas.py", "Independent recalculation + assert vs compute_analysis"),
        ("csv_reader/samples/sample_trades.csv", "Real sample for end-to-end checks"),
    ]:
        ws4.append(list(row))
    _autosize(ws4, [50, 70])

    # Sheet 5 — Legend
    ws5 = wb.create_sheet("Legend")
    ws5.append(["Color / Tag", "Meaning"])
    _style_header(ws5)
    ws5.append(["OK — green", "Matches common industry definition; safe to keep"])
    ws5.append(["REVIEW — yellow", "Works, but definition differs from classic finance formula — confirm with team"])
    ws5.append(["BUG RISK — orange", "Likely incorrect aggregation; fix before trusting the number"])
    ws5["A2"].fill = ok_fill
    ws5["A3"].fill = review_fill
    ws5["A4"].fill = bug_fill
    _autosize(ws5, [22, 80])

    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
